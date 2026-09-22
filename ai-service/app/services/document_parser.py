from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document
from PIL import Image, UnidentifiedImageError

from app.core.config import Settings
from app.schemas.cv import ExtractionMethod
from app.services.normalizer import normalize_text
from app.services.ocr_service import OcrEngine


class CvProcessingError(Exception):
    def __init__(self, status_code: int, detail: str) -> None:
        super().__init__(detail)
        self.status_code = status_code
        self.detail = detail


@dataclass(frozen=True)
class DocumentBlock:
    text: str
    page: int
    bbox: tuple[float, float, float, float]
    font_size: float = 0
    is_bold: bool = False
    column: int = 0


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    media_type: str
    page_count: int | None
    extraction_method: ExtractionMethod
    ocr_applied: bool
    warnings: list[str] = field(default_factory=list)
    blocks: list[DocumentBlock] = field(default_factory=list)
    layout: str = "UNSTRUCTURED"


class DocumentParser:
    _SUPPORTED = {".pdf", ".docx", ".png", ".jpg", ".jpeg"}
    _MEDIA_TYPES = {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
    }

    def __init__(self, settings: Settings, ocr_engine: OcrEngine) -> None:
        self.settings = settings
        self.ocr_engine = ocr_engine

    def parse(self, filename: str, declared_media_type: str | None, data: bytes) -> ExtractedDocument:
        extension = Path(filename or "").suffix.casefold()
        if extension not in self._SUPPORTED:
            raise CvProcessingError(415, "Unsupported CV format; use PDF, DOCX, PNG, JPG, or JPEG")
        if not data:
            raise CvProcessingError(400, "Uploaded CV is empty")
        if len(data) > self.settings.max_upload_size_mb * 1024 * 1024:
            raise CvProcessingError(413, "Uploaded CV exceeds the configured size limit")

        self._validate_signature(extension, data)
        expected_media_type = self._MEDIA_TYPES[extension]
        if declared_media_type and declared_media_type not in {
            expected_media_type,
            "application/octet-stream",
        }:
            raise CvProcessingError(415, "File content type does not match its extension")

        if extension == ".pdf":
            result = self._parse_pdf(data)
        elif extension == ".docx":
            result = self._parse_docx(data)
        else:
            result = self._parse_image(data, expected_media_type)

        normalized_lines = [
            normalized
            for line in result.text.splitlines()
            if (normalized := normalize_text(line))
        ]
        text = "\n".join(normalized_lines)
        if not text:
            raise CvProcessingError(400, "No readable text could be extracted from the CV")
        if len(text) > self.settings.max_extracted_text_chars:
            text = text[: self.settings.max_extracted_text_chars]
            result.warnings.append("EXTRACTED_TEXT_TRUNCATED")
        return ExtractedDocument(
            text=text,
            media_type=result.media_type,
            page_count=result.page_count,
            extraction_method=result.extraction_method,
            ocr_applied=result.ocr_applied,
            warnings=result.warnings,
            blocks=result.blocks,
            layout=result.layout,
        )

    @staticmethod
    def _validate_signature(extension: str, data: bytes) -> None:
        if extension == ".pdf" and not data.startswith(b"%PDF-"):
            raise CvProcessingError(415, "File extension does not match PDF content")
        if extension == ".docx" and not data.startswith(b"PK"):
            raise CvProcessingError(415, "File extension does not match DOCX content")
        if extension == ".png" and not data.startswith(b"\x89PNG\r\n\x1a\n"):
            raise CvProcessingError(415, "File extension does not match PNG content")
        if extension in {".jpg", ".jpeg"} and not data.startswith(b"\xff\xd8\xff"):
            raise CvProcessingError(415, "File extension does not match JPEG content")

    def _parse_pdf(self, data: bytes) -> ExtractedDocument:
        try:
            document = pymupdf.open(stream=data, filetype="pdf")
        except Exception as exc:
            raise CvProcessingError(400, "CV PDF is corrupt or unreadable") from exc
        with document:
            if document.needs_pass:
                raise CvProcessingError(400, "Password-protected CV PDFs are not supported")
            if document.page_count > self.settings.max_pdf_pages:
                raise CvProcessingError(413, "CV PDF exceeds the configured page limit")

            page_texts: list[str] = []
            all_blocks: list[DocumentBlock] = []
            layouts: list[str] = []
            ocr_pages = 0
            for page_number, page in enumerate(document, start=1):
                direct_blocks = self._extract_pdf_blocks(page, page_number)
                direct_text = "\n".join(block.text for block in direct_blocks)
                useful_chars = len(re.sub(r"\W", "", direct_text, flags=re.UNICODE))
                if useful_chars >= self.settings.pdf_text_min_chars_per_page:
                    ordered, layout = self._order_pdf_blocks(direct_blocks, page.rect.width)
                    page_texts.append("\n".join(block.text for block in ordered))
                    all_blocks.extend(ordered)
                    layouts.append(layout)
                    continue
                render_scale = 150 / 72
                rendered_pixels = int(page.rect.width * render_scale) * int(
                    page.rect.height * render_scale
                )
                if rendered_pixels > self.settings.max_image_pixels:
                    raise CvProcessingError(413, "CV PDF page exceeds the OCR pixel limit")
                pixmap = page.get_pixmap(dpi=150, alpha=False)
                image = Image.open(BytesIO(pixmap.tobytes("png"))).convert("RGB")
                page_texts.append(self.ocr_engine.read_text(image))
                ocr_pages += 1
                layouts.append("UNSTRUCTURED")

            if ocr_pages == 0:
                method: ExtractionMethod = "PDF_TEXT"
            elif ocr_pages == document.page_count:
                method = "PDF_OCR"
            else:
                method = "PDF_TEXT_WITH_OCR"
            layout = "MULTI_COLUMN" if "MULTI_COLUMN" in layouts else (
                "SINGLE_COLUMN" if all_blocks else "UNSTRUCTURED"
            )
            warnings = ["MULTI_COLUMN_LAYOUT_DETECTED"] if layout == "MULTI_COLUMN" else []
            return ExtractedDocument(
                text="\n\n".join(page_texts),
                media_type=self._MEDIA_TYPES[".pdf"],
                page_count=document.page_count,
                extraction_method=method,
                ocr_applied=ocr_pages > 0,
                warnings=warnings,
                blocks=all_blocks,
                layout=layout,
            )

    @staticmethod
    def _extract_pdf_blocks(page, page_number: int) -> list[DocumentBlock]:
        results: list[DocumentBlock] = []
        for block in page.get_text("dict", sort=False).get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                spans = line.get("spans", [])
                text = normalize_text("".join(span.get("text", "") for span in spans))
                if not text:
                    continue
                raw_bbox = line.get("bbox", block.get("bbox", (0, 0, 0, 0)))
                results.append(
                    DocumentBlock(
                        text=text,
                        page=page_number,
                        bbox=tuple(float(value) for value in raw_bbox),
                        font_size=max((float(span.get("size", 0)) for span in spans), default=0),
                        is_bold=any(
                            "bold" in str(span.get("font", "")).casefold()
                            or bool(int(span.get("flags", 0)) & 16)
                            for span in spans
                        ),
                    )
                )
        return results

    @staticmethod
    def _order_pdf_blocks(
        blocks: list[DocumentBlock], page_width: float
    ) -> tuple[list[DocumentBlock], str]:
        natural = sorted(blocks, key=lambda item: (item.bbox[1], item.bbox[0]))
        if len(blocks) < 4:
            return natural, "SINGLE_COLUMN"

        x_positions = sorted({round(block.bbox[0], 1) for block in blocks})
        gaps = [(right - left, (left + right) / 2) for left, right in zip(x_positions, x_positions[1:])]
        largest_gap, divider = max(gaps, default=(0, page_width / 2))
        left = [block for block in blocks if (block.bbox[0] + block.bbox[2]) / 2 < divider]
        right = [block for block in blocks if (block.bbox[0] + block.bbox[2]) / 2 >= divider]
        vertical_overlap = bool(left and right) and (
            min(max(item.bbox[1] for item in left), max(item.bbox[1] for item in right))
            > max(min(item.bbox[1] for item in left), min(item.bbox[1] for item in right))
        )
        if not (
            largest_gap >= page_width * 0.18
            and len(left) >= 2
            and len(right) >= 2
            and vertical_overlap
        ):
            return natural, "SINGLE_COLUMN"

        ordered: list[DocumentBlock] = []
        for column, group in enumerate((left, right), start=1):
            ordered.extend(
                DocumentBlock(
                    text=item.text,
                    page=item.page,
                    bbox=item.bbox,
                    font_size=item.font_size,
                    is_bold=item.is_bold,
                    column=column,
                )
                for item in sorted(group, key=lambda item: (item.bbox[1], item.bbox[0]))
            )
        return ordered, "MULTI_COLUMN"

    def _parse_docx(self, data: bytes) -> ExtractedDocument:
        try:
            with zipfile.ZipFile(BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > self.settings.max_docx_entries:
                    raise CvProcessingError(413, "DOCX contains too many archive entries")
                total_size = sum(item.file_size for item in entries)
                if total_size > self.settings.max_docx_uncompressed_mb * 1024 * 1024:
                    raise CvProcessingError(413, "DOCX uncompressed content exceeds the limit")
                if "word/document.xml" not in {item.filename for item in entries}:
                    raise CvProcessingError(415, "File is not a valid DOCX document")
        except CvProcessingError:
            raise
        except (zipfile.BadZipFile, OSError) as exc:
            raise CvProcessingError(400, "CV DOCX is corrupt or unreadable") from exc

        try:
            document = Document(BytesIO(data))
            parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    line = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if line:
                        parts.append(line)
        except Exception as exc:
            raise CvProcessingError(400, "CV DOCX is corrupt or unreadable") from exc
        warnings = [] if parts else ["DOCX_EMBEDDED_IMAGE_OCR_NOT_SUPPORTED"]
        return ExtractedDocument(
            text="\n".join(parts),
            media_type=self._MEDIA_TYPES[".docx"],
            page_count=None,
            extraction_method="DOCX_TEXT",
            ocr_applied=False,
            warnings=warnings,
            layout="FLOW",
        )

    def _parse_image(self, data: bytes, media_type: str) -> ExtractedDocument:
        try:
            with Image.open(BytesIO(data)) as image:
                image.verify()
            with Image.open(BytesIO(data)) as image:
                if image.width * image.height > self.settings.max_image_pixels:
                    raise CvProcessingError(413, "CV image exceeds the configured pixel limit")
                rgb_image = image.convert("RGB")
                text = self.ocr_engine.read_text(rgb_image)
        except CvProcessingError:
            raise
        except (UnidentifiedImageError, OSError, ValueError) as exc:
            raise CvProcessingError(400, "CV image is corrupt or unreadable") from exc
        return ExtractedDocument(
            text=text,
            media_type=media_type,
            page_count=1,
            extraction_method="IMAGE_OCR",
            ocr_applied=True,
            layout="UNSTRUCTURED",
        )

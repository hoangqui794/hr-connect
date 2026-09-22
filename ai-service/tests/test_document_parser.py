from io import BytesIO

import pymupdf
import pytest
from docx import Document
from PIL import Image

from app.core.config import Settings
from app.services.document_parser import CvProcessingError, DocumentParser


class FakeOcrEngine:
    def __init__(self, text: str = "OCR extracted text from CV") -> None:
        self.text = text
        self.calls = 0

    def read_text(self, image: Image.Image) -> str:
        self.calls += 1
        return self.text


def _settings(**overrides) -> Settings:
    return Settings(_env_file=None, **overrides)


def _pdf_bytes(text: str | None = None) -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    if text:
        page.insert_text((72, 72), text)
    data = document.tobytes()
    document.close()
    return data


def _two_column_pdf_bytes() -> bytes:
    document = pymupdf.open()
    page = document.new_page(width=600, height=800)
    page.insert_text((40, 80), "NGUYEN VAN AN", fontsize=20)
    page.insert_text((40, 130), "EDUCATION", fontsize=15)
    page.insert_text((40, 160), "Example University")
    page.insert_text((340, 60), "CAREER OBJECTIVE", fontsize=16)
    page.insert_text((340, 95), "Backend engineer building APIs")
    page.insert_text((340, 145), "TECHNICAL SKILLS", fontsize=16)
    page.insert_text((340, 180), "C#, ASP.NET Core, PostgreSQL")
    data = document.tobytes()
    document.close()
    return data


def _docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("NGUYEN VAN AN")
    document.add_paragraph("Kỹ năng")
    document.add_paragraph("ASP.NET Core, PostgreSQL")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "English"
    table.cell(0, 1).text = "Intermediate"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _png_bytes() -> bytes:
    output = BytesIO()
    Image.new("RGB", (32, 32), "white").save(output, format="PNG")
    return output.getvalue()


def test_pdf_uses_text_layer_without_ocr_and_preserves_lines() -> None:
    ocr = FakeOcrEngine()
    parser = DocumentParser(_settings(pdf_text_min_chars_per_page=5), ocr)

    result = parser.parse("candidate.pdf", "application/pdf", _pdf_bytes("Skills\nPython and Docker"))

    assert result.extraction_method == "PDF_TEXT"
    assert result.ocr_applied is False
    assert ocr.calls == 0
    assert "Skills" in result.text


def test_pdf_reconstructs_two_columns_without_interleaving_lines() -> None:
    parser = DocumentParser(_settings(pdf_text_min_chars_per_page=5), FakeOcrEngine())

    result = parser.parse("two-column.pdf", "application/pdf", _two_column_pdf_bytes())

    lines = result.text.splitlines()
    assert result.layout == "MULTI_COLUMN"
    assert "MULTI_COLUMN_LAYOUT_DETECTED" in result.warnings
    assert lines.index("NGUYEN VAN AN") < lines.index("EDUCATION")
    assert lines.index("Example University") < lines.index("CAREER OBJECTIVE")
    assert lines.index("CAREER OBJECTIVE") < lines.index("TECHNICAL SKILLS")
    assert result.blocks


def test_blank_pdf_falls_back_to_ocr() -> None:
    ocr = FakeOcrEngine("NGUYEN VAN AN\nKỹ năng\nPython")
    parser = DocumentParser(_settings(), ocr)

    result = parser.parse("scan.pdf", "application/pdf", _pdf_bytes())

    assert result.extraction_method == "PDF_OCR"
    assert result.ocr_applied is True
    assert ocr.calls == 1
    assert "Kỹ năng" in result.text


def test_docx_extracts_paragraphs_and_tables() -> None:
    parser = DocumentParser(_settings(), FakeOcrEngine())

    result = parser.parse(
        "candidate.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        _docx_bytes(),
    )

    assert result.extraction_method == "DOCX_TEXT"
    assert "Kỹ năng\nASP.NET Core, PostgreSQL" in result.text
    assert "English | Intermediate" in result.text


def test_image_is_read_through_ocr() -> None:
    ocr = FakeOcrEngine("Image CV text")
    parser = DocumentParser(_settings(), ocr)

    result = parser.parse("candidate.png", "image/png", _png_bytes())

    assert result.extraction_method == "IMAGE_OCR"
    assert result.text == "Image CV text"
    assert ocr.calls == 1


@pytest.mark.parametrize(
    ("filename", "media_type", "data"),
    [
        ("candidate.txt", "text/plain", b"plain text"),
        ("candidate.pdf", "application/pdf", b"not a pdf"),
        ("candidate.png", "image/png", b"not an image"),
    ],
)
def test_rejects_unsupported_or_spoofed_files(filename: str, media_type: str, data: bytes) -> None:
    parser = DocumentParser(_settings(), FakeOcrEngine())

    with pytest.raises(CvProcessingError) as error:
        parser.parse(filename, media_type, data)

    assert error.value.status_code == 415


def test_rejects_pdf_over_page_limit() -> None:
    document = pymupdf.open()
    document.new_page()
    document.new_page()
    data = document.tobytes()
    document.close()
    parser = DocumentParser(_settings(max_pdf_pages=1), FakeOcrEngine())

    with pytest.raises(CvProcessingError) as error:
        parser.parse("candidate.pdf", "application/pdf", data)

    assert error.value.status_code == 413


def test_rejects_oversized_pdf_page_before_ocr_rendering() -> None:
    document = pymupdf.open()
    document.new_page(width=5000, height=5000)
    data = document.tobytes()
    document.close()
    ocr = FakeOcrEngine()
    parser = DocumentParser(_settings(max_image_pixels=1_000_000), ocr)

    with pytest.raises(CvProcessingError) as error:
        parser.parse("oversized.pdf", "application/pdf", data)

    assert error.value.status_code == 413
    assert ocr.calls == 0

"""Turn an uploaded CV into a provider-neutral document plus the plain text used to verify quotes."""

from __future__ import annotations

import base64
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document

_IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}
_MIN_TEXT_CHARS_PER_PAGE = 30


class UnsupportedDocumentError(ValueError):
    pass


@dataclass
class PreparedDocument:
    kind: str
    """"pdf" or "image": the LLM reads the file itself. "text": the LLM reads `text` (DOCX)."""
    media_type: str
    data: bytes
    source_text: str
    """Text read locally (PDF text layer or DOCX). Empty for scans/images: the LLM then reads the pixels and its
    own transcription is the only text available, so the result is flagged for manual review."""
    extraction_method: str
    has_text_layer: bool
    page_count: int | None = None
    image_only_pages: list[int] | None = None
    """1-based pages with no text layer (e.g. a header page saved as an image) in an otherwise text PDF.
    Quotes from those pages can only be checked against the LLM transcription."""

    @property
    def base64_data(self) -> str:
        return base64.standard_b64encode(self.data).decode("ascii")


def prepare_document(filename: str, data: bytes) -> PreparedDocument:
    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        return _prepare_pdf(data)
    if extension == ".docx":
        return _prepare_docx(data)
    if extension in _IMAGE_TYPES:
        return PreparedDocument(
            kind="image", media_type=_IMAGE_TYPES[extension], data=data, source_text="",
            extraction_method="IMAGE_VISION", has_text_layer=False, page_count=1,
        )
    raise UnsupportedDocumentError(f"Unsupported file type '{extension}'. Use PDF, DOCX, PNG, JPEG or WEBP.")


def _prepare_pdf(data: bytes) -> PreparedDocument:
    try:
        with pymupdf.open(stream=data, filetype="pdf") as document:
            page_count = document.page_count
            page_texts = [page.get_text("text") for page in document]
    except Exception as exc:  # corrupt or encrypted file
        raise UnsupportedDocumentError("The PDF could not be opened.") from exc

    text = "\n".join(page_texts)
    image_only = [
        number for number, page_text in enumerate(page_texts, start=1)
        if sum(character.isalnum() for character in page_text) < _MIN_TEXT_CHARS_PER_PAGE
    ]
    has_text_layer = page_count > 0 and len(image_only) < page_count
    return PreparedDocument(
        # The LLM reads the PDF itself, so multi-column layouts and scans both work.
        kind="pdf", media_type="application/pdf", data=data,
        source_text=text if has_text_layer else "",
        extraction_method="PDF_TEXT_AND_VISION" if has_text_layer else "PDF_VISION",
        has_text_layer=has_text_layer, page_count=page_count,
        image_only_pages=image_only if has_text_layer and image_only else None,
    )


def _prepare_docx(data: bytes) -> PreparedDocument:
    try:
        document = Document(BytesIO(data))
    except Exception as exc:
        raise UnsupportedDocumentError("The DOCX could not be opened.") from exc

    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    text = "\n".join(line for line in lines if line.strip())
    return PreparedDocument(
        kind="text", media_type="text/plain", data=text.encode("utf-8"), source_text=text,
        extraction_method="DOCX_TEXT", has_text_layer=True,
    )
